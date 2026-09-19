"""Rebuild the fillable template from the customer's original without changing it."""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[2]
source = next(p for p in (root.parent / 'docs').glob('*.docx') if '20260527' in p.name)
document = Document(source)
nodes = list(document.element.body)
replacements = {7: '{{title}}', 8: '{{owner}}', 21: '{{date}}', 24: '{{summary}}',
                25: '{{taskCaption}}', 27: '{{notice}}', 33: '{{problems}}',
                35: '{{nextFocus}}', 38: '{{otherMatters}}'}
for index, text in replacements.items():
    paragraph = next(p for p in document.paragraphs if p._p is nodes[index])
    style = paragraph.runs[0]._r.rPr if paragraph.runs else None
    from copy import deepcopy
    style = deepcopy(style) if style is not None else None
    paragraph.clear()
    run = paragraph.add_run(text)
    if style is not None:
        run._r.insert(0, style)
document.element.body.remove(nodes[29])
for table in document.tables:
    for row in list(table.rows)[2:]:
        table._tbl.remove(row._tr)
    for cell in table.rows[1].cells:
        cell.text = ''
for tag in ['bookmarkStart', 'bookmarkEnd']:
    for node in document.element.iter(qn('w:' + tag)):
        node.getparent().remove(node)
target = root / 'src/main/resources/templates/work-record.docx'
target.parent.mkdir(parents=True, exist_ok=True)
document.save(target)
print('Prepared', target.name)
