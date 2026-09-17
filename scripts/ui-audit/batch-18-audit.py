"""Budget, approval and account reference audit in a fresh disposable fixture."""
import importlib.util
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import urllib.request

ROOT = Path(__file__).resolve().parents[2]
sys.dont_write_bytecode = True
spec = importlib.util.spec_from_file_location('runner', ROOT / 'scripts/run-review-regression.py')
runner = importlib.util.module_from_spec(spec)
spec.loader.exec_module(runner)
schema = json.loads((ROOT / 'target/ui-audit/state.json').read_text())['schema']
assert re.fullmatch(r'biz_review_test_[0-9a-f]{32}', schema)
env = os.environ.copy()
env['MYSQL_PWD'] = env['SHUANGGAO_TEST_DB_PASSWORD']

def sql(statement):
    result = subprocess.run([runner.executable('mysql'), '--host=127.0.0.1', '--port=3306',
        '--user=' + env.get('SHUANGGAO_TEST_DB_USER', 'root'), '--default-character-set=utf8mb4',
        '--batch', '--skip-column-names', schema], input=statement.encode('utf-8'), env=env, capture_output=True, check=True)
    return result.stdout.decode('utf-8').strip()

def call(path, token=None, body=None):
    headers = {'Content-Type': 'application/json'}
    if token:
        headers['Authorization'] = token
    req = urllib.request.Request('http://127.0.0.1:18080/api' + path,
        data=None if body is None else json.dumps(body).encode('utf-8'), headers=headers)
    with urllib.request.urlopen(req, timeout=15) as response:
        raw = response.read().decode('utf-8')
        try:
            return json.loads(raw)
        except json.JSONDecodeError:
            return raw

import io
import queue
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from openpyxl import Workbook, load_workbook

output = ROOT / 'target/ui-audit/evidence-batch-18'
output.mkdir(parents=True, exist_ok=True)
admin = call('/system/login', body={'user_id':110228,'password':'review-fixture-password'})['token']
uploader = call('/system/login', body={'user_id':1910001,'password':'review-fixture-password'})['token']
results = []

# Identical numeric values with different display formats must retain identical amounts.
book = Workbook()
sheet = book.active
sheet.append(['','资金来源','','合计数','中央投入资金'])
for name in ['五年项目总预算','2026年可使用资金','2026年预算安排','上年结转结余资金','到位数']:
    sheet.append(['',name,'',1.2345,1.2345])
sheet.append([])
sheet.append(['资金支持方向','改革任务','','合计数','1 合成任务'])
for name in ['金额（万元）','商品和服务支出','资本性支出','其他支出','指标名称','项目期满的目标值','目标累计实现情况']:
    sheet.append(['',name,'',1.2345,1.2345])
budgets = []
for month, number_format in [(1,'0.00'),(2,'0.0000')]:
    for row in [2,3,4,5,6,9,10,11,12]:
        sheet.cell(row,5).number_format = number_format
    file = output / ('budget-' + str(month) + '.xlsx')
    book.save(file)
    assert load_workbook(file).active['E2'].value == 1.2345
    boundary = 'batch18-isolated-boundary'
    data = (f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="budget-{month}.xlsx"\r\nContent-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n').encode() + file.read_bytes() + f'\r\n--{boundary}--\r\n'.encode()
    req = urllib.request.Request(f'http://127.0.0.1:18080/api/budget/import?year=2026&month={month}', data=data,
        headers={'Authorization':admin,'Content-Type':'multipart/form-data; boundary='+boundary})
    with urllib.request.urlopen(req,timeout=15) as response:
        imported = json.loads(response.read())
    assert 'sheet' in imported, imported
    budgets.append({'displayFormat':number_format,'cellValue':1.2345,'savedSource':imported['sources'][0]['fiveYearTotal'],'savedTask':imported['tasks'][0]['amount']})
assert budgets[0]['savedSource'] == budgets[0]['savedTask'] == 1.23
assert budgets[1]['savedSource'] == budgets[1]['savedTask'] == 1.2345
results.append({'id':'18-1','samples':budgets})

# Hold the real row write lock until two real HTTP requests have read the pending state.
added = call('/achievement/add', uploader, {'category':1,'level':'省级','achName':'合成并发成果','department':'合成颁发单位','comment':'合成并发测试','gotTime':'2026-09-17 10:00:00','isCompetition':1,'yiDengJiang':1})
assert isinstance(added, str) and '添加成功' in added, added
sub_id = int(sql('SELECT MAX(sub_id) FROM biz_achievement_submission'))
lock = subprocess.Popen([runner.executable('mysql'),'--host=127.0.0.1','--port=3306','--user='+env.get('SHUANGGAO_TEST_DB_USER','root'),
    '--batch','--unbuffered','--skip-column-names',schema],stdin=subprocess.PIPE,stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True,env=env)
ready = queue.Queue()
threading.Thread(target=lambda: [ready.put(line.strip()) for line in lock.stdout],daemon=True).start()
pool = ThreadPoolExecutor(max_workers=2)
try:
    lock.stdin.write(f"START TRANSACTION; SELECT sub_id FROM biz_achievement_submission WHERE sub_id={sub_id} FOR UPDATE; SELECT 'READY';\n")
    lock.stdin.flush()
    while ready.get(timeout=10) != 'READY':
        pass
    futures = [pool.submit(call,'/achievement/audit',admin,{'sub_id':sub_id,'is_pass':decision,'title':'合成并发审核','content':'合成并发审核'}) for decision in [True,False]]
    deadline = time.monotonic()+10
    waiting = 0
    while time.monotonic()<deadline:
        waiting = int(sql("SELECT COUNT(*) FROM information_schema.innodb_trx t JOIN information_schema.processlist p ON t.trx_mysql_thread_id=p.ID WHERE p.DB='"+schema+"' AND t.trx_state='LOCK WAIT'"))
        if waiting == 2:
            break
        time.sleep(0.1)
    assert waiting == 2, 'Both real requests must reach their row updates before release'
    lock.stdin.write('COMMIT;\nquit\n'); lock.stdin.flush()
    responses = [future.result(timeout=15) for future in futures]
finally:
    if lock.poll() is None:
        lock.stdin.close()
        lock.wait(timeout=15)
    pool.shutdown(wait=True)
assert responses == ['成果已归档','成果已退回'], responses
logs = [json.loads(row) for row in sql(f"SELECT JSON_OBJECT('action',action_type,'before',pre_status,'after',post_status) FROM biz_achievement_audit_log WHERE sub_id={sub_id} AND action_type<>'submit' ORDER BY log_id").splitlines()]
assert len(logs) == 2 and all(row['before'] == 10 for row in logs)
notices = sql("SELECT title FROM sys_notice WHERE to_user_id=1910001 ORDER BY notice_id").splitlines()
assert '成果已归档' in notices and '成果已退回' in notices
results.append({'id':'18-2','responses':responses,'logs':logs,'noticeTitles':notices,
    'finalStatus':int(sql(f'SELECT flow_status FROM biz_achievement_submission WHERE sub_id={sub_id}'))})

# Deletion protects task assignees, but ignores a performance-only pending reviewer.
call('/system/users/add',admin,{'userId':918003,'deptId':920001,'userName':'batch18_reviewer','nickName':'合成绩效审核人','password':'review-fixture-password','role':'1','status':'0'})
reviewer_token = call('/system/login',body={'user_id':918003,'password':'review-fixture-password'})['token']
sql("""
INSERT INTO biz_performance (perf_id,project_id,perf_code,perf_name,target_value,data_type,dept_id,principal_id,auditor_id,leader_id)
VALUES (958003,1,'2.18.3','合成删除账号绩效',10,'1',920001,910002,918003,910001);
INSERT INTO biz_performance_year (year_id,perf_id,year,target_value,actual_value) VALUES (958004,958003,2026,10,0);
""")
owner = call('/system/login',body={'user_id':910001,'password':'review-fixture-password'})['token']
submitted = call('/performance/submit?pref_id=958003&actual_value=3&year=2026',owner,{})
assert submitted == '绩效已提交', submitted
deleted = call('/system/users/delete/918003',admin,{})
assert '删除成功' in deleted
assert sql('SELECT is_delete FROM sys_user WHERE user_id=918003') == '1'
assert sql('SELECT current_handler_id FROM biz_performance_submission WHERE perf_id=958003') == '918003'
try:
    call('/performance/audit/todo',reviewer_token)
    raise AssertionError('Deleted reviewer token must be rejected')
except urllib.error.HTTPError as error:
    assert error.code == 401
results.append({'id':'18-3','deleteResponse':deleted,'pendingHandler':918003,'reviewerTokenStatus':401,
    'flowStatus':int(sql('SELECT flow_status FROM biz_performance_submission WHERE perf_id=958003'))})
(output/'results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'reproduced':[row['id'] for row in results]}))
